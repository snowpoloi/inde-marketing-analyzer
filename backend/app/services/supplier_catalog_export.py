from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAX_EXPORT_ROWS = 50_000
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# The main prices are per INDE sale (set), while invoice unit costs are per piece.
PRODUCT_COLUMNS = [
    ("supplier", "Supplier", "text", 24),
    ("name", "Supplier product", "text", 64),
    ("category", "Category", "text", 48),
    ("supplier_code", "Supplier model", "text", 22),
    ("supplier_sku", "Supplier SKU", "text", 22),
    ("ean", "EAN", "text", 22),
    ("opencart_sku", "INDE SKU", "text", 22),
    ("match_method", "Match status", "text", 22),
    ("quantity", "Supplier stock", "number", 16),
    ("sale_quantity", "Pieces / sale", "number", 16),
    ("inde_price", "INDE price (VAT incl.)", "money", 23),
    ("inde_price_net", "INDE price (net)", "money", 23),
    ("aade_cost_sale_net", "AADE cost / sale (net)", "money", 23),
    ("aade_cost_net", "AADE cost / piece (net)", "money", 23),
    ("aade_cost_date", "AADE cost date", "date", 18),
    ("aade_mark", "AADE MARK", "text", 24),
    ("gross_profit_per_sale", "Gross profit / sale (net)", "money", 24),
    ("gross_profit_per_unit", "Gross profit / piece (net)", "money", 24),
    ("gross_margin_percent", "Gross margin %", "percent", 18),
    ("wholesale_price_net", "XML wholesale / sale (net)", "money", 25),
    ("wholesale_price_per_piece_net", "XML wholesale / piece (net)", "money", 25),
    ("retail_price_gross", "XML retail / sale (gross)", "money", 25),
    ("retail_price_per_piece_gross", "XML retail / piece (gross)", "money", 25),
    ("inde_price_basis", "INDE price basis", "text", 24),
    ("margin_status", "Margin status", "text", 28),
    ("last_seen_at", "Last XML sync (UTC)", "text", 28),
]

SUMMARY_COLUMNS = [
    ("supplier", "Supplier", "text", 24),
    ("vat_number", "Supplier AFM", "text", 18),
    ("invoices", "Invoices", "number", 14),
    ("credit_notes", "Credit notes", "number", 14),
    ("purchases_net", "AADE purchases (net)", "money", 24),
    ("credits_net", "Credits (net)", "money", 22),
    ("net_purchases", "Net purchases", "money", 22),
    ("costed_products_net", "Costed products (net)", "money", 24),
    ("priced_units", "Pieces with margin", "number", 22),
    ("costed_units", "Costed pieces", "number", 18),
    ("average_profit_per_unit", "Avg gross profit / piece (net)", "money", 28),
    ("average_margin_percent", "Weighted gross margin %", "percent", 24),
    ("excluded_conflicts", "Excluded invoice conflicts", "number", 25),
]


def _cell(sheet, value, kind):
    if isinstance(value, datetime):
        value = value.astimezone(timezone.utc).isoformat()
    if kind == "date" and isinstance(value, str):
        value = date.fromisoformat(value)
    if kind == "percent" and value is not None:
        value = Decimal(str(value)) / 100
    if kind == "text" and isinstance(value, date):
        value = value.isoformat()
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell = WriteOnlyCell(sheet, value=value)
    if isinstance(value, str):
        # Preserve identifiers/leading zeroes and never interpret supplier text as formulas.
        cell.data_type = "s"
    cell.number_format = {"text": "@", "money": '#,##0.00" EUR"', "number": "#,##0.####",
                          "percent": "0.00%", "date": "dd/mm/yyyy"}[kind]
    return cell


def _sheet(workbook, title, columns, rows):
    sheet = workbook.create_sheet(title)
    sheet.freeze_panes = "A2"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    for index, (_, _, _, width) in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    headers = []
    for _, label, _, _ in columns:
        cell = _cell(sheet, label, "text")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176B59")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        headers.append(cell)
    sheet.row_dimensions[1].height = 34
    sheet.append(headers)
    count = 1
    for row in rows:
        sheet.append([_cell(sheet, row.get(key), kind) for key, _, kind, _ in columns])
        count += 1
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{count}"


def catalog_workbook(products: dict, summary: dict, filters: dict, supplier: str) -> bytes:
    workbook = Workbook(write_only=True)
    _sheet(workbook, "Products", PRODUCT_COLUMNS, products["rows"])
    _sheet(workbook, "Supplier summary", SUMMARY_COLUMNS, summary["rows"])
    metadata = [
        ("Exported at (UTC)", datetime.now(timezone.utc).isoformat()),
        ("Supplier", supplier),
        ("From", summary["date_from"]), ("To", summary["date_to"]),
        ("Search", filters["q"]), ("INDE catalog", filters["match"]),
        ("Supplier stock", filters["availability"]), ("Category", filters["category"] or "All categories"),
        ("Only with gross margin", "Yes" if filters["has_margin"] else "No"),
        ("Sort column", filters["sort_by"]), ("Sort direction", filters["sort_direction"]),
        ("Exported products", products["total"]),
        ("Product prices", "Current catalog prices and latest confirmed AADE unit costs; not filtered by the report period."),
        ("Supplier summary period", "From / To applies to supplier invoice totals. Product filters do not alter these totals."),
        ("Gross margin", "(INDE selling price net - AADE cost per sale net) / INDE selling price net. Excludes freight, ads and other expenses."),
        ("Summary purchases", "Full AADE invoice net totals include invoice freight and other charges. Not proof of payment."),
        ("Summary profit", "Potential profit at current catalog prices, weighted by purchased pieces. Not realized sales profit."),
        ("Missing values", "Blank cells mean unknown, not zero. Identifiers are text; prices are numeric EUR values."),
    ]
    _sheet(workbook, "Selections", [("selection", "Selection", "text", 28), ("value", "Value", "text", 110)],
           ({"selection": label, "value": value} for label, value in metadata))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
