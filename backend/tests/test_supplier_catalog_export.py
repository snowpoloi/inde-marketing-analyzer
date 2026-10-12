from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.models import SupplierCatalogFeed, SupplierCatalogProduct
from app.services.supplier_aade_costs import accept_invoice, invoice_preview
from app.services.supplier_catalog_export import catalog_workbook
from test_supplier_aade_costs import seed, payload
from test_supplier_catalog_analysis import client, purchase


PERIOD = {"date_from": "2026-09-01", "date_to": "2026-09-30"}


def workbook(response):
    assert response.status_code == 200, response.text[:300] if response.status_code != 200 else ""
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert response.headers["content-disposition"] == 'attachment; filename="supplier-catalog-2026-09-01-2026-09-30.xlsx"'
    assert response.headers["cache-control"] == "no-store"
    return load_workbook(BytesIO(response.content))


def records(sheet):
    rows = list(sheet.values)
    return [dict(zip(rows[0], row)) for row in rows[1:]]


def test_export_includes_all_pages_and_keeps_numeric_values_and_leading_zeroes(db):
    user, supplier, _, item, fiscal = seed(db)
    accept_invoice(db, fiscal.id, payload(invoice_preview(db, fiscal.id, supplier.id)), user)
    for n in range(65):
        db.add(SupplierCatalogProduct(feed_id=item.feed_id, supplier_code=f"00{n:03}",
            supplier_sku=f"SKU-{n}", ean="0012345678901", name=f"Product {n:03}", quantity=n,
            is_current=True, last_seen_at=datetime.now(timezone.utc)))
    db.commit()
    with client(db, admin=True) as http:
        filters = {"feed_id": str(item.feed_id), "sort_by": "name", "sort_direction": "desc"}
        expected = http.get("/api/supplier-catalog/products", params={**filters, "limit": 100}).json()
        # Page selection never limits the export; all matching rows are exported.
        book = workbook(http.get("/api/supplier-catalog/export", params={**PERIOD, **filters, "offset": 50, "limit": 50}))
    assert book.sheetnames == ["Products", "Supplier summary", "Selections"]
    rows = records(book["Products"])
    assert len(rows) == 66
    assert [row["Supplier model"] for row in rows] == [row["supplier_code"] for row in expected["rows"]]
    assert rows[0]["EAN"] == "0012345678901"
    assert rows[0]["AADE cost / sale (net)"] is None
    assert rows[0]["Gross margin %"] is None
    assert rows[-1]["Supplier model"] == "0268292"
    assert rows[-1]["INDE price (VAT incl.)"] == 124
    assert rows[-1]["INDE price (net)"] == 100
    assert rows[-1]["AADE cost / sale (net)"] == 70
    assert rows[-1]["Gross profit / sale (net)"] == 30
    assert rows[-1]["Gross margin %"] == 0.3
    assert rows[-1]["AADE cost date"].date() == date(2026, 9, 1)
    assert rows[-1]["AADE MARK"] == "400-test"
    assert book["Products"]["S67"].number_format == "0.00%"
    assert book["Products"].freeze_panes == "A2"
    assert book["Products"].auto_filter.ref == "A1:Z67"
    summary = records(book["Supplier summary"])[0]
    assert summary["AADE purchases (net)"] == 145
    assert summary["Costed products (net)"] == 140
    selections = {row["Selection"]: row["Value"] for row in records(book["Selections"])}
    assert selections["From"] == "2026-09-01" and selections["To"] == "2026-09-30"
    assert selections["Exported products"] == 66


@pytest.mark.parametrize("filters", [
    {"has_margin": "true", "sort_by": "gross_margin_percent", "sort_direction": "asc"},
    {"has_margin": "true", "sort_by": "gross_margin_percent", "sort_direction": "desc"},
    {"match": "unmatched", "availability": "out_of_stock", "category": "Chairs"},
    {"match": "matched", "availability": "in_stock", "q": "zero", "category": "Tables"},
    {"q": "SKU%", "sort_by": "supplier_sku"},
    {"feed_id": str(uuid4())},
])
def test_export_uses_the_same_filters_and_order_as_the_table(db, filters):
    user, supplier, _, item, fiscal = seed(db)
    accept_invoice(db, fiscal.id, payload(invoice_preview(db, fiscal.id, supplier.id)), user)
    _, zero, _ = purchase(db, user, supplier, item.feed_id, "zero", 124, 100, 2, "Zero margin")
    zero.category = "Tables"
    purchase(db, user, supplier, item.feed_id, "negative", 124, 120, 1, "Negative margin")
    feed = SupplierCatalogFeed(code="OTHER", name="Other supplier", adapter="megapap", encrypted_url="test")
    db.add(feed); db.flush()
    for target in (item.feed_id, feed.id):
        db.add(SupplierCatalogProduct(feed_id=target, supplier_code="unknown", supplier_sku="SKU%",
            name="Unknown chair", category="Chairs", quantity=0, is_current=True,
            last_seen_at=datetime.now(timezone.utc)))
    db.commit()
    filters = {"feed_id": str(item.feed_id), **filters}
    with client(db, admin=True) as http:
        expected = http.get("/api/supplier-catalog/products", params={**filters, "limit": 100}).json()
        rows = records(workbook(http.get("/api/supplier-catalog/export", params={**PERIOD, **filters}))["Products"])
    assert [row["Supplier model"] for row in rows] == [row["supplier_code"] for row in expected["rows"]]
    if filters.get("has_margin"):
        assert sorted(row["Gross margin %"] for row in rows) == [-0.2, 0, 0.3]


def test_export_keeps_set_prices_and_period_summary_separate(db):
    user, supplier, _, item, _ = seed(db)
    product, set_item, _ = purchase(db, user, supplier, item.feed_id, "SET4", 248, 30, 8, "Set")
    set_item.details = {"minimum": "4"}
    set_item.wholesale_price_net = 30; set_item.retail_price_gross = 62
    db.commit()
    with client(db, admin=True) as http:
        params = {"date_from": "2026-08-01", "date_to": "2026-08-31", "q": "SET4"}
        response = http.get("/api/supplier-catalog/export", params=params)
        assert response.status_code == 200
        book = load_workbook(BytesIO(response.content))
    row = records(book["Products"])[0]
    assert row["Pieces / sale"] == 4 and row["INDE price (net)"] == 200
    assert row["AADE cost / piece (net)"] == 30 and row["AADE cost / sale (net)"] == 120
    assert row["Gross profit / sale (net)"] == 80 and row["Gross profit / piece (net)"] == 20
    assert row["XML wholesale / sale (net)"] == 120 and row["XML retail / sale (gross)"] == 248
    assert records(book["Supplier summary"])[0]["AADE purchases (net)"] == 0
    db.refresh(product); assert product.price == 248


def test_export_validation_access_and_row_limit(db, monkeypatch):
    from app.api.deps import get_current_user
    from app.api.routes import supplier_catalog
    seed(db)
    with client(db, admin=True) as http:
        for filters in [{"sort_by": "bad"}, {"has_margin": "bad"}, {"match": "bad"},
                        {"availability": "bad"}, {"sort_direction": "bad"}, {"q": "x" * 201},
                        {"date_from": "2026-10-01"}, {"date_to": "invalid"}]:
            assert http.get("/api/supplier-catalog/export", params={**PERIOD, **filters}).status_code == 422
        assert http.get("/api/supplier-catalog/export").status_code == 422
        monkeypatch.setattr(supplier_catalog, "MAX_EXPORT_ROWS", 0)
        response = http.get("/api/supplier-catalog/export", params=PERIOD)
        assert response.status_code == 413 and "narrower filters" in response.json()["detail"]
    with client(db, admin=False) as http:
        assert http.get("/api/supplier-catalog/export", params=PERIOD).status_code == 403
        del http.app.dependency_overrides[get_current_user]
        assert http.get("/api/supplier-catalog/export", params=PERIOD).status_code == 401


def test_workbook_treats_untrusted_text_as_text_and_keeps_unknowns_blank():
    products = {"rows": [{"name": '=HYPERLINK("https://example.test", "bad")', "supplier_code": "0012",
        "supplier_sku": "+123", "ean": "@formula", "supplier": "Bad\x01 supplier",
        "aade_cost_sale_net": None, "gross_profit_per_sale": Decimal("-3.5")}], "total": 1}
    filters = {"q": "=1+1", "match": "all", "availability": "all", "category": "", "has_margin": False,
               "sort_by": "name", "sort_direction": "asc"}
    book = load_workbook(BytesIO(catalog_workbook(products, {"rows": [], **PERIOD}, filters, "=1+1")))
    for sheet in book:
        assert all(cell.data_type != "f" for row in sheet for cell in row)
    assert book["Products"]["B2"].value == products["rows"][0]["name"]
    assert book["Products"]["D2"].value == "0012" and book["Products"]["D2"].data_type == "s"
    assert book["Products"]["A2"].value == "Bad supplier"
    assert book["Products"]["M2"].value is None
    assert book["Products"]["Q2"].value == -3.5
