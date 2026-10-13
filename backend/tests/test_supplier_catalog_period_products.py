from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models import AADEDocument, SupplierCatalogFeed, SupplierCatalogProduct
from app.services.supplier_aade_costs import accept_invoice, invoice_preview
from app.services.supplier_catalog_service import refresh_state
from app.services.supplier_catalog_settings import queue_cost_refresh
from test_supplier_aade_costs import seed, payload, auto_policy
from test_supplier_catalog_analysis import catalog, client, purchase
from test_supplier_catalog_export import records


def costed(db):
    user, supplier, own, item, fiscal = seed(db)
    accept_invoice(db, fiscal.id, payload(invoice_preview(db, fiscal.id, supplier.id)), user)
    purchase(db, user, supplier, item.feed_id, "SECOND", 248, 150, 1, "Second")
    return user, supplier, own, item, fiscal


@pytest.mark.parametrize("has_margin", [False, True])
@pytest.mark.parametrize("sort_by", ["name", "aade_cost_sale_net", "gross_margin_percent"])
def test_period_scopes_products_counts_and_costs_before_pagination(db, has_margin, sort_by):
    _, _, _, _, _ = costed(db)
    options = dict(product_scope="period", date_from=date(2026, 9, 2), date_to=date(2026, 9, 2),
                   has_margin=has_margin, sort_by=sort_by)
    result = catalog(db, **options)
    assert result["total"] == result["summary"]["products"] == result["summary"]["matched"] == 1
    assert result["rows"][0]["supplier_code"] == "SECOND"
    assert result["rows"][0]["aade_cost_date"] == date(2026, 9, 2)
    assert result["rows"][0]["aade_cost_net"] == 150
    assert catalog(db, **options, offset=1)["rows"] == []
    assert catalog(db, **options, q="0268292")["total"] == 0


def test_period_does_not_fall_back_to_an_old_cost_or_use_credit_note_as_cost(db):
    _, supplier, _, _, _ = costed(db)
    db.add(AADEDocument(source_endpoint="RequestDocs", identity_key="credit-only", mark="credit-only",
        issuer_vat=supplier.vat_number, counterpart_vat="802216736", issue_date=date(2026, 10, 13),
        aa="3402", series="C", invoice_type="5.2", currency="EUR", document_direction="expense",
        net_value=Decimal("631.86"), raw={"record_type":"full_document"}))
    db.commit()
    result = catalog(db, product_scope="period", date_from=date(2026, 10, 13), date_to=date(2026, 10, 13), has_margin=True)
    assert result["rows"] == [] and result["total"] == result["summary"]["products"] == 0
    assert catalog(db, has_margin=True)["total"] == 2


def test_period_cost_uses_latest_inside_period_not_a_later_price(db):
    user, supplier, own, item, fiscal = seed(db)
    accept_invoice(db, fiscal.id, payload(invoice_preview(db, fiscal.id, supplier.id)), user)
    newer = AADEDocument(source_endpoint="RequestDocs", identity_key="newer", mark="newer",
        issuer_vat=supplier.vat_number, counterpart_vat=fiscal.counterpart_vat, issue_date=date(2026, 9, 3),
        aa="newer", series="A", invoice_type="1.1", currency="EUR", document_direction="expense",
        net_value=80, vat_amount=Decimal("19.2"), gross_value=Decimal("99.2"),
        raw={"record_type":"full_document", "invoiceDetails":[{"lineNumber":1,"itemCode":item.supplier_code,
            "quantity":1,"measurementUnit":1,"netValue":80,"vatAmount":19.2}]})
    db.add(newer); db.commit()
    accept_invoice(db, newer.id, payload(invoice_preview(db, newer.id, supplier.id)), user)
    period = catalog(db, product_scope="period", date_from=date(2026, 9, 1), date_to=date(2026, 9, 1))["rows"][0]
    assert period["aade_cost_net"] == 70 and period["gross_margin_percent"] == 30
    assert catalog(db)["rows"][0]["aade_cost_net"] == 80
    fiscal.is_cancelled = True; db.commit()
    assert catalog(db, product_scope="period", date_from=date(2026, 9, 1), date_to=date(2026, 9, 1))["total"] == 0


def test_cost_from_one_supplier_does_not_include_same_product_at_another_supplier(db):
    _, _, own, _, _ = costed(db)
    feed = SupplierCatalogFeed(code="OTHER", name="Other", adapter="megapap", encrypted_url="fixture")
    db.add(feed); db.flush()
    db.add(SupplierCatalogProduct(feed_id=feed.id, supplier_code="OTHER", name="Other", product_catalog_id=own.id,
        is_current=True, last_seen_at=datetime.now(timezone.utc)))
    db.commit()
    options = dict(product_scope="period", date_from=date(2026, 9, 1), date_to=date(2026, 9, 1))
    assert catalog(db, **options)["total"] == 1
    assert catalog(db, **options, feed_id=feed.id)["total"] == 0


@pytest.mark.parametrize("day,expected", [("2026-09-01", ["0268292"]), ("2026-09-02", ["SECOND"]), ("2026-10-13", [])])
def test_export_and_table_share_the_same_period_including_empty_days(db, day, expected):
    costed(db)
    params = dict(product_scope="period", date_from=day, date_to=day, has_margin="true")
    with client(db, admin=True) as http:
        response = http.get("/api/supplier-catalog/products", params=params)
        assert response.status_code == 200, response.text
        result = response.json()
        response = http.get("/api/supplier-catalog/export", params=params)
        assert response.status_code == 200, response.text[:500]
        book = load_workbook(BytesIO(response.content))
    assert [row["supplier_code"] for row in result["rows"]] == expected
    assert [row["Supplier model"] for row in records(book["Products"])] == expected
    selections = {row["Selection"]:row["Value"] for row in records(book["Selections"])}
    assert selections["Exported products"] == len(expected)
    assert selections["Product scope"] == "AADE costs in period"


@pytest.mark.parametrize("params", [{"product_scope":"bad"}, {"product_scope":"period"},
    {"product_scope":"period","date_from":"2026-09-01"},
    {"product_scope":"period","date_from":"2026-09-02","date_to":"2026-09-01"}])
def test_period_validation_is_shared_by_export_and_products(db, params):
    seed(db)
    with client(db, admin=True) as http:
        assert http.get("/api/supplier-catalog/products", params=params).status_code == 422
        assert http.get("/api/supplier-catalog/export", params=params).status_code == 422


def test_refresh_state_changes_for_sync_and_costs_without_returning_secrets(db):
    user, supplier, _, item, fiscal = seed(db)
    auto_policy(db, user, supplier)
    before = refresh_state(db)
    assert before["cost_revisions"][str(item.feed_id)] is None
    accept_invoice(db, fiscal.id, payload(invoice_preview(db, fiscal.id, supplier.id)), user)
    after = refresh_state(db)
    assert after["cost_revisions"][str(item.feed_id)] is not None
    queue_cost_refresh(db); db.commit()
    assert refresh_state(db)["sync_revision"] != before["sync_revision"]
    assert "encrypted_url" not in after["feeds"][0] and "url" not in after["feeds"][0]
    with client(db, admin=True) as http:
        assert http.get("/api/supplier-catalog/refresh-state").status_code == 200
    with client(db, admin=False) as http:
        assert http.get("/api/supplier-catalog/refresh-state").status_code == 403
